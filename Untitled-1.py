"""Coletor de notícias (versão de teste)
 
O que faz:
1. Lê feeds RSS de portais de notícias (não precisa de chave de API)
2. Extrai título, resumo, link, data, fonte e categoria de cada notícia
3. Remove duplicadas (pelo link)
4. Salva tudo em um arquivo noticias.xlsx na mesma pasta do script
 
Instalação (uma vez):
    pip install requests openpyxl
 
Uso:
    python coletor_noticias.py
"""
 
import re
import html
import xml.etree.ElementTree as ET
from datetime import datetime
from email.utils import parsedate_to_datetime
from pathlib import Path
 
import requests
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
 
# ---------------------------------------------------------------------------
# CONFIGURAÇÃO
# ---------------------------------------------------------------------------
 
# (fonte, categoria, url do feed RSS)
# Se algum feed sair do ar, é só trocar a URL ou comentar a linha.
FEEDS = [
    ("G1", "Geral", "https://g1.globo.com/rss/g1/"),
    ("G1", "Economia", "https://g1.globo.com/rss/g1/economia/"),
    ("G1", "Tecnologia", "https://g1.globo.com/rss/g1/tecnologia/"),
    ("Agência Brasil", "Geral", "https://agenciabrasil.ebc.com.br/rss/ultimasnoticias/feed.xml"),
]
 
# Máximo de notícias por feed (para o teste não ficar gigante)
LIMITE_POR_FEED = 20
 
# Arquivo de saída: mesma pasta do script
ARQUIVO_SAIDA = Path(__file__).parent / "noticias.xlsx"
 
HEADERS = {"User-Agent": "Mozilla/5.0 (coletor-noticias-teste)"}
TIMEOUT = 15  # segundos
 
 
# ---------------------------------------------------------------------------
# FUNÇÕES AUXILIARES
# ---------------------------------------------------------------------------
 
def limpar_html(texto):
    """Tira tags HTML e entidades (&amp; etc.) do resumo."""
    if not texto:
        return ""
    texto = re.sub(r"<[^>]+>", " ", texto)   # remove tags
    texto = html.unescape(texto)              # &amp; -> &
    texto = re.sub(r"\s+", " ", texto)        # espaços duplicados
    return texto.strip()
 
 
def converter_data(texto):
    """Converte a data do RSS (formato RFC 822) em datetime. Retorna None se falhar."""
    if not texto:
        return None
    try:
        dt = parsedate_to_datetime(texto)
        return dt.replace(tzinfo=None)  # Excel não aceita fuso horário
    except (TypeError, ValueError):
        return None
 
 
def ler_feed(fonte, categoria, url):
    """Baixa um feed RSS e devolve uma lista de dicionários (uma notícia por item)."""
    print(f"Lendo {fonte} / {categoria} ...")
    resposta = requests.get(url, headers=HEADERS, timeout=TIMEOUT)
    resposta.raise_for_status()
 
    raiz = ET.fromstring(resposta.content)
    noticias = []
 
    # No RSS, cada notícia fica em <channel><item>...</item></channel>
    for item in raiz.iter("item"):
        noticias.append({
            "titulo": limpar_html(item.findtext("title")),
            "resumo": limpar_html(item.findtext("description")),
            "link": (item.findtext("link") or "").strip(),
            "data_publicacao": converter_data(item.findtext("pubDate")),
            "fonte": fonte,
            "categoria": categoria,
        })
        if len(noticias) >= LIMITE_POR_FEED:
            break
 
    print(f"  -> {len(noticias)} notícias")
    return noticias
 
 
def coletar_tudo():
    """Percorre todos os feeds, junta as notícias e remove duplicadas pelo link."""
    todas = []
    vistos = set()
 
    for fonte, categoria, url in FEEDS:
        try:
            noticias = ler_feed(fonte, categoria, url)
        except Exception as erro:
            # Se um feed falhar, segue para o próximo
            print(f"  !! Erro em {fonte}/{categoria}: {erro}")
            continue
 
        for n in noticias:
            if not n["link"] or n["link"] in vistos:
                continue
            vistos.add(n["link"])
            todas.append(n)
 
    # Mais recentes primeiro (as sem data vão para o fim)
    todas.sort(key=lambda n: n["data_publicacao"] or datetime.min, reverse=True)
    return todas
 
 
def salvar_xlsx(noticias, caminho):
    """Grava a lista de notícias em um arquivo Excel."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Noticias"
 
    colunas = [
        ("Título", "titulo", 60),
        ("Resumo", "resumo", 80),
        ("Link", "link", 50),
        ("Data publicação", "data_publicacao", 20),
        ("Fonte", "fonte", 18),
        ("Categoria", "categoria", 15),
    ]
 
    # Cabeçalho
    for i, (nome, _, largura) in enumerate(colunas, start=1):
        c = ws.cell(row=1, column=i, value=nome)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[c.column_letter].width = largura
 
    # Linhas
    for linha, n in enumerate(noticias, start=2):
        for i, (_, chave, _) in enumerate(colunas, start=1):
            c = ws.cell(row=linha, column=i, value=n[chave])
            c.alignment = Alignment(vertical="top", wrap_text=(chave in ("titulo", "resumo")))
            if chave == "data_publicacao" and n[chave]:
                c.number_format = "dd/mm/yyyy hh:mm"
 
    ws.freeze_panes = "A2"                       # congela o cabeçalho
    ws.auto_filter.ref = ws.dimensions           # filtros nas colunas
 
    wb.save(caminho)
 
 
# ---------------------------------------------------------------------------
# EXECUÇÃO
# ---------------------------------------------------------------------------
 
if __name__ == "__main__":
    noticias = coletar_tudo()
 
    if not noticias:
        print("Nenhuma notícia coletada. Verifique a internet ou as URLs dos feeds.")
    else:
        salvar_xlsx(noticias, ARQUIVO_SAIDA)
        print(f"\nPronto! {len(noticias)} notícias salvas em: {ARQUIVO_SAIDA}")